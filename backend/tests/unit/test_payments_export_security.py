# pyright: reportUnknownParameterType=false, reportMissingParameterType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownVariableType=false, reportUnusedCallResult=false, reportOptionalMemberAccess=false, reportOptionalSubscript=false
"""Regression tests for the payments workbook export.

``GET /api/v1/payments/export/excel`` previously had no role guard and no
actor scoping: any authenticated user (including a ``requester``) could
download every payment in the system, including bank references and notes.
These tests lock down the fixed behaviour.
"""

from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    PaymentRecord,
    PaymentStatus,
    POItem,
    POStatus,
    PurchaseOrder,
    PurchaseRequisition,
    Supplier,
    User,
)
from app.services import export_excel as svc

from .test_export_ledger_api import _auth_headers, _client, _user


async def _payment_fixture(db, *, actor: User, department_id) -> tuple[PurchaseOrder, POItem]:
    supplier = Supplier(code=f"SUP-PAY-{uuid4().hex[:6].upper()}", name="Payments Export Supplier")
    db.add(supplier)
    await db.flush()

    pr = PurchaseRequisition(
        pr_number=f"PR-PAY-{uuid4().hex[:6].upper()}",
        title="Payments export",
        status="approved",
        requester_id=actor.id,
        company_id=actor.company_id,
        department_id=department_id,
        currency="CNY",
        total_amount=Decimal("300.00"),
    )
    db.add(pr)
    await db.flush()

    po = PurchaseOrder(
        po_number=f"PO-PAY-{uuid4().hex[:6].upper()}",
        pr_id=pr.id,
        supplier_id=supplier.id,
        company_id=actor.company_id,
        status=POStatus.CONFIRMED.value,
        currency="CNY",
        total_amount=Decimal("300.00"),
        amount_paid=Decimal("0"),
        created_by_id=actor.id,
    )
    db.add(po)
    await db.flush()

    item = POItem(
        po_id=po.id,
        line_no=1,
        item_name="Payments export item",
        qty=Decimal("1"),
        uom="EA",
        unit_price=Decimal("300.00"),
        amount=Decimal("300.00"),
    )
    db.add(item)
    await db.flush()

    db.add(
        PaymentRecord(
            payment_number=f"PAY-{uuid4().hex[:8].upper()}",
            po_id=po.id,
            installment_no=1,
            amount=Decimal("300.00"),
            currency="CNY",
            status=PaymentStatus.CONFIRMED.value,
            payment_method="wire",
            transaction_ref="SECRET-REF",
            notes="SECRET-NOTE",
        )
    )
    await db.flush()
    return po, item


def _payment_rows(content: bytes) -> list[list[object]]:
    worksheet = load_workbook(BytesIO(content))["Payments"]
    return [
        list(row)
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row[0] not in (None, "", "合计 Total")
    ]


async def test_payments_export_denies_dept_manager(seeded_db_session):
    db = seeded_db_session
    bob = await _user(db, "bob")

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/payments/export/excel",
            headers=_auth_headers(bob),
        )

    assert resp.status_code == 403
    assert resp.json()["code"] == "insufficient_role"


async def test_payments_export_requires_authentication(seeded_db_session):
    db = seeded_db_session

    async with _client(db) as client:
        resp = await client.get("/api/v1/payments/export/excel")

    assert resp.status_code == 401


async def test_render_payments_xlsx_scopes_rows_to_actor(seeded_db_session):
    db = seeded_db_session
    alice = await _user(db)
    bob = await _user(db, "bob")

    own_po, _ = await _payment_fixture(db, actor=alice, department_id=bob.department_id)
    hidden_po, _ = await _payment_fixture(db, actor=alice, department_id=None)

    as_dept_manager = _payment_rows(await svc.render_payments_xlsx(db, actor=bob))
    visible_numbers = {row[1] for row in as_dept_manager}
    assert own_po.po_number in visible_numbers
    assert hidden_po.po_number not in visible_numbers
    assert all(row[8] in (None, "") for row in as_dept_manager)  # payment_method denied
    assert all(row[9] in (None, "") for row in as_dept_manager)  # transaction_ref denied
    assert all(row[11] in (None, "") for row in as_dept_manager)  # notes denied

    as_auditor = _payment_rows(await svc.render_payments_xlsx(db, actor=alice))
    auditor_numbers = {row[1] for row in as_auditor}
    assert {own_po.po_number, hidden_po.po_number} <= auditor_numbers


async def test_render_payments_xlsx_gates_sensitive_columns(seeded_db_session):
    from app.models import User as UserModel

    db = seeded_db_session
    alice = await _user(db)
    admin = (await db.execute(select(UserModel).where(UserModel.username == "admin"))).scalar_one()
    po, _ = await _payment_fixture(db, actor=alice, department_id=alice.department_id)

    it_buyer_rows = [
        row
        for row in _payment_rows(await svc.render_payments_xlsx(db, actor=alice))
        if row[1] == po.po_number
    ]
    assert len(it_buyer_rows) == 1
    # it_buyer may not read payment_method / transaction_ref / notes
    assert it_buyer_rows[0][8] in (None, "")
    assert it_buyer_rows[0][9] in (None, "")
    assert it_buyer_rows[0][11] in (None, "")
    assert it_buyer_rows[0][4] == 300.0  # amount is still allowed

    admin_rows = [
        row
        for row in _payment_rows(await svc.render_payments_xlsx(db, actor=admin))
        if row[1] == po.po_number
    ]
    assert len(admin_rows) == 1
    assert admin_rows[0][8] == "wire"
    assert admin_rows[0][9] == "SECRET-REF"
    assert admin_rows[0][11] == "SECRET-NOTE"


async def test_render_payments_xlsx_without_actor_keeps_legacy_behaviour(seeded_db_session):
    """Internal callers that do not pass an actor keep the raw rows."""
    db = seeded_db_session
    alice = await _user(db)
    po, _ = await _payment_fixture(db, actor=alice, department_id=alice.department_id)

    rows = [
        row for row in _payment_rows(await svc.render_payments_xlsx(db)) if row[1] == po.po_number
    ]
    assert len(rows) == 1
    assert rows[0][8] == "wire"
    assert rows[0][9] == "SECRET-REF"
    assert rows[0][11] == "SECRET-NOTE"
