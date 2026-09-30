# pyright: reportUnknownParameterType=false, reportMissingParameterType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownVariableType=false, reportUnusedCallResult=false, reportOptionalMemberAccess=false, reportOptionalSubscript=false
"""Service-level tests for the generic export framework and its datasets."""

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.field_authz import FIELD_PERMISSIONS
from app.models import (
    Contract,
    DeliveryPlan,
    Invoice,
    InvoiceLine,
    InvoiceStatus,
    Item,
    PaymentRecord,
    POItem,
    POStatus,
    ProcurementCategory,
    PurchaseOrder,
    PurchaseRequisition,
    Shipment,
    ShipmentItem,
    ShipmentStatus,
    SKUPriceAnomaly,
    SKUPriceBenchmark,
    SKUPriceRecord,
    Supplier,
    User,
)
from app.services import export_datasets  # noqa: F401  (registers the datasets)
from app.services import export_registry as registry
from app.services.export_excel import _FORMULA_PREFIXES


def _suffix() -> str:
    return uuid4().hex[:8].upper()


async def _user(db, username: str = "alice") -> User:
    return (await db.execute(select(User).where(User.username == username))).scalar_one()


async def _build_world(db, actor: User) -> dict[str, str]:
    """Create one row (or row graph) for every dataset and return markers."""
    supplier = Supplier(
        code=f"SUP-X-{_suffix()}",
        name=f"Export Supplier {_suffix()}",
        contact_name="Contact",
        contact_phone="123",
        contact_email="c@example.com",
        tax_number="TAX-1",
        payee_name="Payee",
        payee_bank="Bank",
        payee_bank_account="6222",
    )
    category = ProcurementCategory(
        code=f"CAT-{_suffix()}", label_zh="测试分类", label_en="Test Category"
    )
    db.add(category)
    await db.flush()
    item = Item(
        code=f"ITM-{_suffix()}",
        name=f"Export Item {_suffix()}",
        category="HPC",
        category_id=category.id,
        uom="EA",
        specification="spec",
    )
    db.add_all([supplier, item])
    await db.flush()

    pr = PurchaseRequisition(
        pr_number=f"PR-X-{_suffix()}",
        title=f"Export PR {_suffix()}",
        status="approved",
        requester_id=actor.id,
        company_id=actor.company_id,
        department_id=actor.department_id,
        currency="CNY",
        total_amount=Decimal("1000.00"),
    )
    db.add(pr)
    await db.flush()

    from app.models import PRItem

    db.add(
        PRItem(
            pr_id=pr.id,
            line_no=1,
            item_id=item.id,
            item_name=item.name,
            qty=Decimal("2"),
            uom="EA",
            unit_price=Decimal("500"),
            amount=Decimal("1000"),
        )
    )
    await db.flush()

    po = PurchaseOrder(
        po_number=f"PO-X-{_suffix()}",
        pr_id=pr.id,
        supplier_id=supplier.id,
        company_id=actor.company_id,
        status=POStatus.CONFIRMED.value,
        currency="CNY",
        total_amount=Decimal("1000.00"),
        amount_paid=Decimal("0"),
        created_by_id=actor.id,
    )
    db.add(po)
    await db.flush()

    po_item = POItem(
        po_id=po.id,
        line_no=1,
        item_id=item.id,
        item_name=item.name,
        qty=Decimal("2"),
        uom="EA",
        unit_price=Decimal("500"),
        amount=Decimal("1000"),
        qty_received=Decimal("1"),
    )
    db.add(po_item)
    await db.flush()

    contract = Contract(
        contract_number=f"CT-X-{_suffix()}",
        po_id=po.id,
        supplier_id=supplier.id,
        title=f"Export Contract {_suffix()}",
        currency="CNY",
        total_amount=Decimal("1000.00"),
        signed_date=date(2026, 1, 5),
        notes="contract note",
    )
    db.add(contract)

    plan = DeliveryPlan(
        po_id=po.id,
        item_id=item.id,
        plan_name=f"Plan {_suffix()}",
        planned_qty=2,
        planned_date=date(2026, 3, 1),
        status="planned",
        created_by_id=actor.id,
    )
    db.add(plan)

    shipment = Shipment(
        shipment_number=f"SH-X-{_suffix()}",
        po_id=po.id,
        batch_no=1,
        status=ShipmentStatus.ARRIVED.value,
        carrier="DHL",
        tracking_number="TRK-1",
        expected_date=date(2026, 3, 1),
        actual_date=date(2026, 3, 2),
        received_by_id=actor.id,
    )
    db.add(shipment)
    await db.flush()
    db.add(
        ShipmentItem(
            shipment_id=shipment.id,
            po_item_id=po_item.id,
            line_no=1,
            item_name=item.name,
            qty_shipped=Decimal("2"),
            qty_received=Decimal("1"),
            unit_price=Decimal("500"),
        )
    )

    invoice = Invoice(
        internal_number=f"INV-X-{_suffix()}",
        invoice_number=f"FP-X-{_suffix()}",
        supplier_id=supplier.id,
        invoice_date=date(2026, 3, 5),
        subtotal=Decimal("1000"),
        tax_amount=Decimal("130"),
        total_amount=Decimal("1130"),
        currency="CNY",
        status=InvoiceStatus.MATCHED.value,
        is_fully_matched=True,
        notes="invoice note",
    )
    db.add(invoice)
    await db.flush()
    db.add(
        InvoiceLine(
            invoice_id=invoice.id,
            po_item_id=po_item.id,
            line_no=1,
            item_name=item.name,
            qty=Decimal("2"),
            unit_price=Decimal("500"),
            subtotal=Decimal("1000"),
            tax_amount=Decimal("130"),
        )
    )

    payment = PaymentRecord(
        payment_number=f"PAY-X-{_suffix()}",
        po_id=po.id,
        contract_id=contract.id,
        installment_no=1,
        amount=Decimal("1000.00"),
        currency="CNY",
        status="confirmed",
        due_date=date(2026, 4, 1),
        payment_date=date(2026, 4, 2),
        payment_method="bank_transfer",
        transaction_ref="TRX-X-1",
        notes="payment note",
    )
    db.add(payment)

    db.add(
        SKUPriceRecord(
            item_id=item.id,
            supplier_id=supplier.id,
            price=Decimal("499.5"),
            currency="CNY",
            quotation_date=date(2026, 3, 3),
            source_type="manual",
            source_ref="REF-1",
            entered_by_id=actor.id,
            notes="price note",
        )
    )
    db.add(
        SKUPriceBenchmark(
            item_id=item.id,
            window_days=90,
            avg_price=Decimal("500"),
            median_price=Decimal("500"),
            stddev=Decimal("5"),
            min_price=Decimal("480"),
            max_price=Decimal("520"),
            sample_size=7,
        )
    )
    db.add(
        SKUPriceAnomaly(
            item_id=item.id,
            baseline_avg_price=Decimal("500"),
            observed_price=Decimal("700"),
            deviation_pct=Decimal("40"),
            severity="warning",
            status="new",
            notes="anomaly note",
        )
    )
    await db.flush()

    return {
        "supplier": supplier.name,
        "supplier_code": supplier.code,
        "item_name": item.name,
        "category_id": str(category.id),
        "item_code": item.code,
        "pr_title": pr.title,
        "contract_title": contract.title,
        "plan_name": plan.plan_name,
        "item": item.code,
        "pr": pr.pr_number,
        "po": po.po_number,
        "contract": contract.contract_number,
        "shipment": shipment.shipment_number,
        "invoice": invoice.internal_number,
        "sku": item.name,
        "ledger": po.po_number,
        "payment": payment.payment_number,
    }


# dataset -> [(sheet index, column header, marker key)]: the marker must appear
# in that exact cell, which catches a column getter reading the wrong index.
COLUMN_CHECKS: dict[str, list[tuple[int, str, str]]] = {
    "purchase_requisitions": [
        (0, "申请单号 PR No.", "pr"),
        (0, "标题 Title", "pr_title"),
        (1, "申请单号 PR No.", "pr"),
    ],
    "contracts": [
        (0, "合同编号 Contract No.", "contract"),
        (1, "合同编号 Contract No.", "contract"),
    ],
    "delivery_plans": [(0, "采购订单号 PO No.", "po")],
    "shipments": [(0, "到货单号 Shipment No.", "shipment"), (0, "物料 Item", "item_name")],
    "invoices": [(0, "内部编号 Internal No.", "invoice"), (1, "内部编号 Internal No.", "invoice")],
    "sku_prices": [(0, "物料 Item", "sku"), (1, "物料 Item", "sku"), (2, "物料 Item", "sku")],
    "suppliers": [(0, "供应商编码 Code", "supplier_code"), (0, "名称 Name", "supplier")],
    "items": [(0, "物料编码 Code", "item"), (0, "名称 Name", "item_name")],
    "procurement_ledger": [(0, "采购订单号 PO No.", "po")],
    "payments": [(0, "付款单号 Payment No.", "payment"), (0, "采购订单号 PO No.", "po")],
}

MARKERS = {
    "purchase_requisitions": "pr",
    "contracts": "contract",
    "delivery_plans": "po",
    "shipments": "shipment",
    "invoices": "invoice",
    "sku_prices": "sku",
    "suppliers": "supplier",
    "items": "item",
    "procurement_ledger": "ledger",
    "payments": "payment",
}


def _no_filters(dataset) -> registry.ExportFilters:
    """Only use filters the dataset declares, so nothing is rejected."""
    return registry.ExportFilters()


async def test_every_registered_dataset_renders_xlsx_and_csv(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    assert len(registry.REGISTRY) >= 9

    for key, dataset in registry.REGISTRY.items():
        filters = _no_filters(dataset)
        payload, filename = await registry.run_export(
            db,
            dataset_key=key,
            actor=actor,
            filters=filters,
            export_format="xlsx",
            max_rows=5000,
        )
        assert payload.startswith(b"PK"), key
        assert filename.startswith(f"mica-{key}-"), key

        workbook = load_workbook(BytesIO(payload))
        assert workbook.sheetnames == [sheet.title for sheet in dataset.sheets], key
        for sheet in dataset.sheets:
            assert [cell.value for cell in workbook[sheet.title][1]] == sheet.headers, key

        # Column 0 must be a 1-based row number, and the next column must not
        # be the placeholder: a regression here shifts every value right.
        first_sheet_rows = list(
            workbook[dataset.sheets[0].title].iter_rows(min_row=2, values_only=True)
        )
        data_rows = [r for r in first_sheet_rows if r and r[0] is not None]
        assert data_rows, f"{key}: no data rows"
        assert [r[0] for r in data_rows] == list(range(1, len(data_rows) + 1)), key
        assert not any(isinstance(r[1], int) and r[1] == 0 for r in data_rows), key

        # The marker proves the loader returned the row we created.
        marker = markers[MARKERS[key]]
        first_sheet = workbook[dataset.sheets[0].title]
        flat = [str(value) for row in first_sheet.iter_rows(values_only=True) for value in row]
        assert any(marker in value for value in flat), f"{key}: marker {marker} missing"

        # Values must land in the right column, not merely somewhere in the row.
        for sheet_index, header, marker_key in COLUMN_CHECKS[key]:
            sheet = workbook[dataset.sheets[sheet_index].title]
            headers = [cell.value for cell in sheet[1]]
            assert header in headers, f"{key}: missing header {header}"
            column = headers.index(header)
            values = [
                row[column]
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if row[0] is not None
            ]
            expected = markers[marker_key]
            assert any(str(value) == expected for value in values), (
                f"{key} sheet{sheet_index}: {header} expected {expected}, got {values}"
            )

        csv_payload, csv_name = await registry.run_export(
            db,
            dataset_key=key,
            actor=actor,
            filters=filters,
            export_format="csv",
            max_rows=5000,
        )
        assert csv_payload.startswith(b"\xef\xbb\xbf"), key
        assert csv_name.endswith(".csv"), key
        assert dataset.sheets[0].headers[0] in csv_payload.decode("utf-8-sig")


async def test_sku_sheets_do_not_mix_their_rows(seeded_db_session):
    """Benchmark rows must not be emitted into the anomaly sheet (and vice versa)."""
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    payload, _ = await registry.run_export(
        db,
        dataset_key="sku_prices",
        actor=actor,
        filters=registry.ExportFilters(),
        export_format="xlsx",
        max_rows=5000,
    )
    workbook = load_workbook(BytesIO(payload))

    bench = workbook["行情基准 Price Benchmarks"]
    bench_rows = [row for row in bench.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert bench_rows, "expected at least one benchmark row"
    assert all(isinstance(row[3], int) for row in bench_rows), bench_rows  # window_days

    anomalies = workbook["价格异常 Price Anomalies"]
    anomaly_rows = [
        row for row in anomalies.iter_rows(min_row=2, values_only=True) if row[0] is not None
    ]
    assert anomaly_rows, "expected at least one anomaly row"
    assert all(row[6] in {"info", "warning", "critical"} for row in anomaly_rows), (
        anomaly_rows
    )  # severity, not a price


async def test_datasets_load_without_help_from_the_identity_map(seeded_db_session):
    """A fresh request session must not need the identity map.

    Production runs with ``expire_on_commit=True`` and a per-request session, so
    a relationship that is used but not eagerly loaded raises MissingGreenlet
    (HTTP 500). Expunging the identity map reproduces that here, which the
    same-session tests above cannot.
    """
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    db.expunge_all()

    for key in registry.REGISTRY:
        payload, _ = await registry.run_export(
            db,
            dataset_key=key,
            actor=actor,
            filters=registry.ExportFilters(),
            export_format="xlsx",
            max_rows=5000,
        )
        assert payload.startswith(b"PK"), key


async def test_catalog_lists_datasets_and_is_role_aware(seeded_db_session):
    catalog = registry.catalog("admin")
    keys = {entry["key"] for entry in catalog}
    # A role outside EXPORT_ROLES gets nothing, and every export role gets the
    # full set (no dataset silently hides itself from one of them).
    assert registry.catalog("requester") == []
    assert registry.catalog("dept_manager") == []
    for role in sorted(registry.EXPORT_ROLES):
        assert {entry["key"] for entry in registry.catalog(role)} == keys, role
    # ``datasets_for`` and the per-dataset ``roles`` field are the mechanism.
    assert {ds.key for ds in registry.datasets_for("admin")} == keys
    assert {
        "purchase_requisitions",
        "contracts",
        "delivery_plans",
        "shipments",
        "invoices",
        "sku_prices",
        "suppliers",
        "items",
        "procurement_ledger",
    } <= keys

    for entry in catalog:
        assert entry["label_zh"] and entry["label_en"]
        assert entry["sheets"], entry["key"]
        for sheet in entry["sheets"]:
            assert sheet["columns"], entry["key"]
            for column in sheet["columns"]:
                assert column["key"] and column["label_zh"] and column["label_en"]


async def test_unknown_dataset_is_not_found(seeded_db_session):
    with pytest.raises(HTTPException) as excinfo:
        registry.get_dataset("nope")
    assert excinfo.value.status_code == 404
    assert excinfo.value.detail == "export.dataset_not_found"


async def test_unsupported_filter_is_rejected(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    # suppliers declares only the keyword filter
    with pytest.raises(HTTPException) as excinfo:
        await registry.run_export(
            db,
            dataset_key="suppliers",
            actor=actor,
            filters=registry.ExportFilters(supplier_id=uuid4()),
            export_format="xlsx",
            max_rows=5000,
        )
    assert excinfo.value.status_code == 400
    assert excinfo.value.detail == "export.unsupported_filter"


async def test_role_outside_export_roles_is_refused(seeded_db_session):
    db = seeded_db_session
    bob = await _user(db, "bob")

    with pytest.raises(HTTPException) as excinfo:
        await registry.run_export(
            db,
            dataset_key="items",
            actor=bob,
            filters=registry.ExportFilters(),
            export_format="xlsx",
            max_rows=5000,
        )
    assert excinfo.value.status_code == 403


async def test_row_cap_is_enforced(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    with pytest.raises(HTTPException) as excinfo:
        await registry.run_export(
            db,
            dataset_key="items",
            actor=actor,
            filters=registry.ExportFilters(),
            export_format="xlsx",
            max_rows=0,
        )
    assert excinfo.value.detail == "export.too_many_rows"


async def test_registering_unknown_cerbos_kind_fails_loudly():
    """A dataset whose gated kind is missing from FIELD_PERMISSIONS would fail
    open, so registration must refuse it."""
    sheet = registry.ExportSheetSpec(
        title_zh="X",
        title_en="X",
        columns=[
            registry.ExportColumn(
                key="secret",
                header_zh="秘密",
                header_en="Secret",
                get=lambda row: row[0],
                field="not_a_real_kind.secret",
            )
        ],
    )
    dataset = registry.ExportDataset(
        key=f"bogus-{_suffix()}",
        label_zh="X",
        label_en="X",
        filters=frozenset(),
        sheets=[sheet],
        loader=None,  # type: ignore[arg-type]
    )
    with pytest.raises(RuntimeError, match="fail open"):
        registry.register(dataset)


async def test_field_gating_blanks_restricted_columns(seeded_db_session):
    db = seeded_db_session
    alice = await _user(db)
    markers = await _build_world(db, alice)
    supplier_code = (
        await db.execute(select(Supplier.code).where(Supplier.name == markers["supplier"]))
    ).scalar_one()

    payload, _ = await registry.run_export(
        db,
        dataset_key="suppliers",
        actor=alice,
        filters=registry.ExportFilters(),
        export_format="xlsx",
        max_rows=5000,
    )
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook[registry.get_dataset("suppliers").sheets[0].title]
    headers = [cell.value for cell in sheet[1]]
    code_column = headers.index("供应商编码 Code")
    # The export is unfiltered, so pick our own row rather than the first one.
    row = next(
        [cell.value for cell in candidate]
        for candidate in sheet.iter_rows(min_row=2)
        if candidate[code_column].value == supplier_code
    )
    by_header = dict(zip(headers, row, strict=True))

    # it_buyer may not read settlement data
    # openpyxl reads a blank string cell back as None
    for header in ("税号 Tax No.", "银行账号 Bank Account", "开户行 Bank", "收款户名 Payee"):
        assert by_header[header] in (None, ""), header
    # ...but does see the rest
    assert by_header["供应商编码 Code"]
    assert by_header["联系人 Contact"] == "Contact"


async def test_formula_like_text_is_neutralized_in_every_dataset(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    supplier = (
        await db.execute(select(Supplier).where(Supplier.name == markers["supplier"]))
    ).scalar_one()
    supplier.name = "=cmd|'/C calc'!A0"
    supplier.payee_bank = "@SUM(1)"
    contract = (
        await db.execute(select(Contract).where(Contract.contract_number == markers["contract"]))
    ).scalar_one()
    contract.title = "=1+1"
    await db.flush()

    for key in ("suppliers", "contracts", "procurement_ledger", "shipments"):
        payload, _ = await registry.run_export(
            db,
            dataset_key=key,
            actor=actor,
            filters=registry.ExportFilters(),
            export_format="xlsx",
            max_rows=5000,
        )
        workbook = load_workbook(BytesIO(payload))
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", f"{key}!{sheet.title}!{cell.coordinate}"


async def test_csv_prefixes_only_dangerous_text(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    supplier = (
        await db.execute(select(Supplier).where(Supplier.name == markers["supplier"]))
    ).scalar_one()
    supplier.name = "=1+1"
    await db.flush()

    payload, _ = await registry.run_export(
        db,
        dataset_key="suppliers",
        actor=actor,
        filters=registry.ExportFilters(),
        export_format="csv",
        max_rows=5000,
    )
    text = payload.decode("utf-8-sig")
    assert "'=1+1" in text
    assert f"'{supplier.code}" not in text


@pytest.mark.parametrize("dangerous", _FORMULA_PREFIXES)
def test_csv_safe_covers_all_dangerous_prefixes(dangerous):
    from app.services.export_excel import _csv_safe

    assert _csv_safe(f"{dangerous}x").startswith("'")
    assert _csv_safe("safe") == "safe"
    assert _csv_safe(12.5) == 12.5


async def test_date_range_filter_is_inclusive_and_overflow_safe(seeded_db_session):
    """Boundaries are whole UTC days, inclusive, and date.max must not overflow."""
    db = seeded_db_session
    actor = await _user(db)

    supplier = Supplier(code=f"SUP-DR-{_suffix()}", name="Date Range Supplier")
    db.add(supplier)
    await db.flush()
    pr = PurchaseRequisition(
        pr_number=f"PR-DR-{_suffix()}",
        title="Date range",
        status="approved",
        requester_id=actor.id,
        company_id=actor.company_id,
        department_id=actor.department_id,
        currency="CNY",
        total_amount=Decimal("10"),
    )
    db.add(pr)
    await db.flush()
    po = PurchaseOrder(
        po_number=f"PO-DR-{_suffix()}",
        pr_id=pr.id,
        supplier_id=supplier.id,
        company_id=actor.company_id,
        status=POStatus.CONFIRMED.value,
        currency="CNY",
        total_amount=Decimal("10"),
        amount_paid=Decimal("0"),
        created_by_id=actor.id,
    )
    db.add(po)
    await db.flush()
    for number, moment in (
        ("PAY-DR-IN", datetime(2026, 5, 10, 23, 30, tzinfo=UTC)),
        ("PAY-DR-OUT", datetime(2026, 5, 11, 0, 30, tzinfo=UTC)),
    ):
        db.add(
            PaymentRecord(
                payment_number=number,
                po_id=po.id,
                installment_no=1,
                amount=Decimal("5"),
                currency="CNY",
                status="pending",
                created_at=moment,
            )
        )
    await db.flush()

    def _numbers(payload: bytes) -> list[str]:
        sheet = load_workbook(BytesIO(payload))["付款记录 Payments"]
        column = [cell.value for cell in sheet[1]].index("付款单号 Payment No.") + 1
        return [
            sheet.cell(row=row, column=column).value
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=column).value
        ]

    day = registry.ExportFilters(date_from=date(2026, 5, 10), date_to=date(2026, 5, 10))
    payload, _ = await registry.run_export(
        db,
        dataset_key="payments",
        actor=actor,
        filters=day,
        export_format="xlsx",
        max_rows=5000,
    )
    numbers = _numbers(payload)
    assert "PAY-DR-IN" in numbers  # 23:30Z on the 10th is inside that UTC day
    assert "PAY-DR-OUT" not in numbers  # 00:30Z on the 11th is not

    everything = registry.ExportFilters(date_from=date(1, 1, 1), date_to=date(9999, 12, 31))
    payload, _ = await registry.run_export(
        db,
        dataset_key="payments",
        actor=actor,
        filters=everything,
        export_format="xlsx",
        max_rows=5000,
    )
    assert {"PAY-DR-IN", "PAY-DR-OUT"} <= set(_numbers(payload))


async def test_ledger_dataset_matches_legacy_endpoint_shape(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    dataset = registry.get_dataset("procurement_ledger")
    assert [sheet.title_zh for sheet in dataset.sheets] == ["采购台账", "付款明细", "发票明细"]

    payload, _ = await registry.run_export(
        db,
        dataset_key="procurement_ledger",
        actor=actor,
        filters=registry.ExportFilters(),
        export_format="xlsx",
        max_rows=5000,
    )
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ["采购台账 Ledger", "付款明细 Payments", "发票明细 Invoice Lines"]


async def test_every_gated_kind_has_a_permission_entry():
    kinds: set[str] = set()
    for dataset in registry.REGISTRY.values():
        kinds |= set(dataset.gated_fields())
    missing = sorted(k for k in kinds if k not in FIELD_PERMISSIONS)
    assert missing == [], f"kinds missing FIELD_PERMISSIONS: {missing}"
    assert kinds, "expected at least one gated kind"


async def test_status_values_come_from_the_enums(seeded_db_session):
    """The catalog must not let the UI hardcode (and drift from) status lists."""
    from app.models import (
        ContractStatus,
        DeliveryPlanStatus,
        InvoiceStatus,
        PaymentStatus,
        POStatus,
        PRStatus,
        ShipmentStatus,
    )

    expected = {
        "purchase_requisitions": PRStatus,
        "contracts": ContractStatus,
        "delivery_plans": DeliveryPlanStatus,
        "shipments": ShipmentStatus,
        "invoices": InvoiceStatus,
        "payments": PaymentStatus,
        "procurement_ledger": POStatus,
    }

    for entry in registry.catalog("admin"):
        if "status" not in entry["filters"]:
            assert entry["status_values"] == [], entry["key"]
            continue
        enum = expected[entry["key"]]
        assert entry["status_values"] == [m.value for m in enum], entry["key"]

    # catched real drift before: PRStatus.cancelled and InvoiceStatus.mismatched
    assert (
        "cancelled"
        in next(e for e in registry.catalog("admin") if e["key"] == "purchase_requisitions")[
            "status_values"
        ]
    )
    assert (
        "mismatched"
        in next(e for e in registry.catalog("admin") if e["key"] == "invoices")["status_values"]
    )


async def test_dataset_declaring_status_without_values_fails_registration():
    dataset = registry.ExportDataset(
        key=f"nostatus-{_suffix()}",
        label_zh="X",
        label_en="X",
        filters=frozenset({"status"}),
        sheets=[
            registry.ExportSheetSpec(
                title_zh="X",
                title_en="X",
                columns=[
                    registry.ExportColumn(
                        key="a", header_zh="A", header_en="A", get=lambda row: row[0]
                    )
                ],
            )
        ],
        loader=None,  # type: ignore[arg-type]
    )
    with pytest.raises(RuntimeError, match="status_values"):
        registry.register(dataset)
