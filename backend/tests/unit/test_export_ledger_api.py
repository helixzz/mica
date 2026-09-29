# pyright: reportUnknownParameterType=false, reportMissingParameterType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownVariableType=false, reportUnusedCallResult=false, reportOptionalMemberAccess=false, reportOptionalSubscript=false
"""API-level tests for ``GET /api/v1/purchase-orders/export/ledger``.

These exercise the real router (auth guard, query parsing, response headers)
against the service layer, using the savepoint-rolled-back ``seeded_db_session``
so nothing leaks between tests.
"""

from contextlib import asynccontextmanager
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.security import create_access_token
from app.db import get_db
from app.main import app
from app.models import (
    POItem,
    POStatus,
    PurchaseOrder,
    PurchaseRequisition,
    Supplier,
    User,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _auth_headers(user: User) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(subject=str(user.id))}"}


@asynccontextmanager
async def _client(db):
    """HTTP client bound to *db*, restoring the dependency override on exit.

    Leaving ``app.dependency_overrides`` populated leaks a closed session into
    every later test module (the legacy ``seeded_client`` fixture inherits it).
    """

    async def _override_get_db():
        yield db

    app.dependency_overrides[get_db] = _override_get_db
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            yield client
    finally:
        app.dependency_overrides.pop(get_db, None)


async def _user(db, username: str = "alice") -> User:
    return (await db.execute(select(User).where(User.username == username))).scalar_one()


async def _ledger_fixture(
    db, actor: User, *, status: str = POStatus.CONFIRMED.value
) -> tuple[PurchaseOrder, POItem]:
    supplier = Supplier(code=f"SUP-API-{uuid4().hex[:6].upper()}", name="Ledger API Supplier")
    db.add(supplier)
    await db.flush()

    pr = PurchaseRequisition(
        pr_number=f"PR-API-{uuid4().hex[:6].upper()}",
        title="Ledger API",
        status="approved",
        requester_id=actor.id,
        company_id=actor.company_id,
        department_id=actor.department_id,
        currency="CNY",
        total_amount=Decimal("300.00"),
    )
    db.add(pr)
    await db.flush()

    po = PurchaseOrder(
        po_number=f"PO-API-{uuid4().hex[:6].upper()}",
        pr_id=pr.id,
        supplier_id=supplier.id,
        company_id=actor.company_id,
        status=status,
        currency="CNY",
        total_amount=Decimal("300.00"),
        amount_paid=Decimal("0"),
        created_by_id=actor.id,
    )
    db.add(po)
    await db.flush()

    item = POItem(
        po_id=po.id,
        line_no=1,
        item_name="API item",
        qty=Decimal("3"),
        uom="EA",
        unit_price=Decimal("100"),
        amount=Decimal("300.00"),
    )
    db.add(item)
    await db.flush()
    return po, item


async def test_export_ledger_returns_xlsx_attachment(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    po, _ = await _ledger_fixture(db, actor)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"supplier_id": str(po.supplier_id)},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_CONTENT_TYPE
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="mica-procurement-ledger-')
    assert disposition.endswith('.xlsx"')

    workbook = load_workbook(BytesIO(resp.content))
    assert workbook.sheetnames == ["采购台账 Ledger", "付款明细 Payments", "发票明细 Invoice Lines"]
    ledger = workbook["采购台账 Ledger"]
    po_numbers = [ledger.cell(row=row, column=12).value for row in range(2, ledger.max_row + 1)]
    assert po.po_number in po_numbers


async def test_export_ledger_returns_csv(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    po, _ = await _ledger_fixture(db, actor)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"format": "csv", "supplier_id": str(po.supplier_id)},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["content-disposition"].endswith('.csv"')
    assert resp.content.startswith(b"\xef\xbb\xbf")
    assert po.po_number in resp.content.decode("utf-8-sig")


async def test_export_ledger_rejects_unknown_status(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"status": "not-a-status"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 400
    assert resp.json()["code"] == "export.invalid_status"


async def test_export_ledger_rejects_unknown_format(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"format": "pdf"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 422


async def test_export_ledger_denies_dept_manager(seeded_db_session):
    db = seeded_db_session
    bob = await _user(db, "bob")

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            headers=_auth_headers(bob),
        )

    assert resp.status_code == 403
    assert resp.json()["code"] == "insufficient_role"


async def test_export_ledger_requires_authentication(seeded_db_session):
    db = seeded_db_session

    async with _client(db) as client:
        resp = await client.get("/api/v1/purchase-orders/export/ledger")

    assert resp.status_code == 401


async def test_export_ledger_survives_po_id_route_collision(seeded_db_session):
    """``/purchase-orders/{po_id}`` must not swallow the literal export path."""
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200


@pytest.mark.parametrize("statuses", [["confirmed"], ["confirmed", "closed"]])
async def test_export_ledger_accepts_repeated_status_params(seeded_db_session, statuses):
    db = seeded_db_session
    actor = await _user(db)
    kept_po, _ = await _ledger_fixture(db, actor, status=POStatus.CONFIRMED.value)
    closed_po, _ = await _ledger_fixture(db, actor, status=POStatus.CLOSED.value)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params=[("format", "csv"), *[("status", s) for s in statuses]],
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200
    body = resp.content.decode("utf-8-sig")
    assert kept_po.po_number in body
    if "closed" in statuses:
        assert closed_po.po_number in body
    else:
        assert closed_po.po_number not in body


async def test_export_ledger_rejects_inverted_date_range(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"date_from": "2026-06-02", "date_to": "2026-06-01"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 400
    assert resp.json()["code"] == "export.invalid_date_range"


async def test_export_ledger_accepts_maximum_date(seeded_db_session):
    """``date.max`` must not overflow the upper bound into an HTTP 500."""
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            params={"date_from": "0001-01-01", "date_to": "9999-12-31"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200


async def test_export_ledger_reads_max_rows_from_system_parameters(seeded_db_session):
    """The cap must come from system_parameters, not a hardcoded constant."""
    from app.models import SystemParameter
    from app.services.system_params import system_params

    db = seeded_db_session
    actor = await _user(db)
    await _ledger_fixture(db, actor)
    await _ledger_fixture(db, actor)

    param = (
        await db.execute(select(SystemParameter).where(SystemParameter.key == "export.max_rows"))
    ).scalar_one()

    original = param.value
    try:
        param.value = 1
        await db.flush()
        system_params.invalidate()

        async with _client(db) as client:
            resp = await client.get(
                "/api/v1/purchase-orders/export/ledger",
                headers=_auth_headers(actor),
            )

        assert resp.status_code == 400
        assert resp.json()["code"] == "export.too_many_rows"
    finally:
        param.value = original
        await db.flush()
        system_params.invalidate()
