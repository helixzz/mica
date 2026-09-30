# pyright: reportUnknownParameterType=false, reportMissingParameterType=false, reportUnknownMemberType=false, reportUnknownArgumentType=false, reportUnknownVariableType=false, reportUnusedCallResult=false, reportOptionalMemberAccess=false, reportOptionalSubscript=false
"""API-level tests for the generic export endpoints."""

from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook

from .test_export_ledger_api import _auth_headers, _client, _user
from .test_export_registry import _build_world

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def test_export_catalog_lists_datasets_for_the_role(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get("/api/v1/exports", headers=_auth_headers(actor))

    assert resp.status_code == 200
    datasets = resp.json()["datasets"]
    keys = {entry["key"] for entry in datasets}
    assert "purchase_requisitions" in keys
    assert "procurement_ledger" in keys

    ledger = next(entry for entry in datasets if entry["key"] == "procurement_ledger")
    assert ledger["label_zh"] == "采购台账"
    assert "supplier" in ledger["filters"]
    assert len(ledger["sheets"]) == 3
    assert ledger["sheets"][0]["columns"][0]["key"] == "seq"


async def test_export_dataset_returns_xlsx_attachment(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/suppliers",
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_CONTENT_TYPE
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="mica-suppliers-')
    assert disposition.endswith('.xlsx"')

    workbook = load_workbook(BytesIO(resp.content))
    assert workbook.sheetnames == ["供应商 Suppliers"]
    values = [
        cell
        for row in workbook["供应商 Suppliers"].iter_rows(min_row=2, values_only=True)
        for cell in row
    ]
    assert any(markers["supplier"] == value for value in values)


async def test_export_dataset_returns_csv_with_bom(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/items",
            params={"format": "csv"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["content-disposition"].endswith('.csv"')
    assert resp.content.startswith(b"\xef\xbb\xbf")
    assert resp.content.decode("utf-8-sig").startswith("序号 No.,")


async def test_export_dataset_unknown_key_is_404(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/not-a-dataset",
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 404
    assert resp.json()["code"] == "export.dataset_not_found"


async def test_export_dataset_rejects_unsupported_filter(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    # items declares only category + keyword
    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/items",
            params={"supplier_id": "00000000-0000-0000-0000-000000000001"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 400
    assert resp.json()["code"] == "export.unsupported_filter"


async def test_export_dataset_rejects_inverted_date_range(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/contracts",
            params={"date_from": "2026-06-02", "date_to": "2026-06-01"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 400
    assert resp.json()["code"] == "export.invalid_date_range"


async def test_export_dataset_accepts_maximum_date(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/contracts",
            params={"date_from": "0001-01-01", "date_to": "9999-12-31"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200


async def test_export_dataset_rejects_unknown_format(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/items",
            params={"format": "pdf"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 422


async def test_export_dataset_denies_role_outside_export_roles(seeded_db_session):
    db = seeded_db_session
    bob = await _user(db, "bob")

    async with _client(db) as client:
        catalog_resp = await client.get("/api/v1/exports", headers=_auth_headers(bob))
        assert catalog_resp.status_code == 200
        assert catalog_resp.json()["datasets"] == []

        resp = await client.get("/api/v1/exports/items", headers=_auth_headers(bob))

    assert resp.status_code == 403
    assert resp.json()["code"] == "insufficient_role"


async def test_export_dataset_requires_authentication(seeded_db_session):
    db = seeded_db_session

    async with _client(db) as client:
        catalog_resp = await client.get("/api/v1/exports")
        dataset_resp = await client.get("/api/v1/exports/items")

    assert catalog_resp.status_code == 401
    assert dataset_resp.status_code == 401


async def test_export_dataset_applies_declared_filters(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    supplier_id = None
    from sqlalchemy import select as sa_select

    from app.models import Supplier

    supplier_id = (
        await db.execute(sa_select(Supplier.id).where(Supplier.name == markers["supplier"]))
    ).scalar_one()

    async with _client(db) as client:
        kept = await client.get(
            "/api/v1/exports/suppliers",
            params={"q": markers["supplier"]},
            headers=_auth_headers(actor),
        )
        dropped = await client.get(
            "/api/v1/exports/suppliers",
            params={"q": "no-such-supplier-xyz"},
            headers=_auth_headers(actor),
        )
        by_status = await client.get(
            "/api/v1/exports/contracts",
            params=[("status", "active"), ("supplier_id", str(supplier_id))],
            headers=_auth_headers(actor),
        )

    assert kept.status_code == 200
    assert dropped.status_code == 200
    assert by_status.status_code == 200

    # XLSX is a zip, so assert on the parsed workbook rather than raw bytes.
    def _codes(payload: bytes) -> list[str]:
        workbook = load_workbook(BytesIO(payload))
        return [
            row[1]
            for row in workbook["供应商 Suppliers"].iter_rows(min_row=2, values_only=True)
            if row[1]
        ]

    assert len(_codes(kept.content)) == 1
    assert markers["supplier"] in "".join(
        str(cell)
        for row in load_workbook(BytesIO(kept.content))["供应商 Suppliers"].iter_rows(
            values_only=True
        )
        for cell in row
    )
    assert _codes(dropped.content) == []

    # status=active + supplier filter must keep this contract and drop others
    contracts = load_workbook(BytesIO(by_status.content))
    rows = [
        row for row in contracts["合同 Contracts"].iter_rows(min_row=2, values_only=True) if row[1]
    ]
    assert [row[1] for row in rows] == [markers["contract"]]

    async with _client(db) as client:
        other_supplier = await client.get(
            "/api/v1/exports/contracts",
            params=[
                ("status", "active"),
                ("supplier_id", "00000000-0000-0000-0000-0000000000ff"),
            ],
            headers=_auth_headers(actor),
        )

    assert other_supplier.status_code == 200
    assert [
        row
        for row in load_workbook(BytesIO(other_supplier.content))["合同 Contracts"].iter_rows(
            min_row=2, values_only=True
        )
        if row[1]
    ] == []


async def test_ledger_endpoint_still_works_alongside_the_registry(seeded_db_session):
    """The v1.52.0 endpoint must keep behaving after the refactor."""
    db = seeded_db_session
    actor = await _user(db)
    await _build_world(db, actor)

    async with _client(db) as client:
        legacy = await client.get(
            "/api/v1/purchase-orders/export/ledger",
            headers=_auth_headers(actor),
        )
        generic = await client.get(
            "/api/v1/exports/procurement_ledger",
            headers=_auth_headers(actor),
        )

    assert legacy.status_code == 200
    assert generic.status_code == 200
    assert (
        load_workbook(BytesIO(legacy.content)).sheetnames
        == load_workbook(BytesIO(generic.content)).sheetnames
    )


async def test_export_dataset_rejects_unknown_status_value(seeded_db_session):
    """The generic endpoint validates status against the dataset's own values."""
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/invoices",
            params={"status": "not-a-real-status"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 400
    assert resp.json()["code"] == "export.invalid_status"


async def test_export_dataset_accepts_declared_status_value(seeded_db_session):
    db = seeded_db_session
    actor = await _user(db)

    async with _client(db) as client:
        resp = await client.get(
            "/api/v1/exports/invoices",
            params={"status": "mismatched"},
            headers=_auth_headers(actor),
        )

    assert resp.status_code == 200


async def test_role_outside_export_roles_cannot_enumerate_dataset_keys(seeded_db_session):
    """A disallowed role gets 403 even for a key that does not exist."""
    db = seeded_db_session
    bob = await _user(db, "bob")

    async with _client(db) as client:
        existing = await client.get("/api/v1/exports/items", headers=_auth_headers(bob))
        missing = await client.get(
            "/api/v1/exports/definitely-not-a-dataset", headers=_auth_headers(bob)
        )

    assert existing.status_code == 403
    assert missing.status_code == 403
    assert existing.json()["code"] == missing.json()["code"] == "insufficient_role"


async def test_sku_secondary_sheets_respect_item_dimension_filters(seeded_db_session):
    """A category filter must scope the benchmark and anomaly sheets too."""
    from decimal import Decimal

    from app.models import Item, SKUPriceAnomaly, SKUPriceBenchmark
    from app.services import export_registry as registry

    db = seeded_db_session
    actor = await _user(db)
    markers = await _build_world(db, actor)

    other = Item(
        code=f"ITM-OTHER-{markers['item_code']}",
        name="Other category item",
        category="OTHER",
    )
    db.add(other)
    await db.flush()
    db.add(
        SKUPriceBenchmark(
            item_id=other.id,
            window_days=90,
            avg_price=Decimal("1"),
            median_price=Decimal("1"),
            stddev=Decimal("0"),
            min_price=Decimal("1"),
            max_price=Decimal("1"),
            sample_size=1,
        )
    )
    db.add(
        SKUPriceAnomaly(
            item_id=other.id,
            baseline_avg_price=Decimal("1"),
            observed_price=Decimal("2"),
            deviation_pct=Decimal("100"),
            severity="warning",
            status="new",
        )
    )
    await db.flush()

    entry = next(item for item in registry.catalog("admin") if item["key"] == "sku_prices")
    assert "category" in entry["filters"]

    payload, _ = await registry.run_export(
        db,
        dataset_key="sku_prices",
        actor=actor,
        filters=registry.ExportFilters(category_id=UUID(markers["category_id"])),
        export_format="xlsx",
        max_rows=5000,
    )
    workbook = load_workbook(BytesIO(payload))
    sheet_titles = [
        "价格记录 Price Records",
        "行情基准 Price Benchmarks",
        "价格异常 Price Anomalies",
    ]
    # The out-of-category item must appear in NO sheet, including the aggregates.
    for title in sheet_titles:
        values = {
            str(value)
            for row in workbook[title].iter_rows(min_row=2, values_only=True)
            for value in row
        }
        assert other.code not in values, title
    in_category_codes = {
        str(value)
        for row in workbook["价格记录 Price Records"].iter_rows(min_row=2, values_only=True)
        for value in row
    }
    assert markers["item_code"] in in_category_codes
